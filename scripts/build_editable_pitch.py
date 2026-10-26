"""Build SecClaw's pitch with editable text and shapes in every slide."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "slide-repairs"
OUTPUT = ROOT / "docs" / "SecClaw-editable-pitch.pptx"
P = Presentation()
P.slide_width, P.slide_height = Inches(10), Inches(5.625)
BLANK = P.slide_layouts[6]
NAVY, TEAL, GRAY = RGBColor(23, 34, 55), RGBColor(17, 121, 111), RGBColor(83, 96, 117)
BG, WHITE, BORDER = RGBColor(245, 247, 251), RGBColor(255, 255, 255), RGBColor(214, 222, 232)


def slide():
    s = P.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(s, value, x, y, w, h, size=18, color=NAVY, bold=False, margin=0):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Noto Sans"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_before = Pt(0)
        p.space_after = Pt(0)
    return sh


def rect(s, x, y, w, h, fill=WHITE, line=BORDER, rounded=False):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1)
    return sh


def head(s, kicker, title, subtitle=None, title_size=27):
    text(s, kicker, .48, .27, 9.1, .3, 11, TEAL, True)
    title_h = .75 if len(title) < 55 else 1.05
    text(s, title, .48, .58, 9.0, title_h, title_size, NAVY, True)
    sep = .58 + title_h + .09
    rect(s, .48, sep, 1.35, .025, TEAL, TEAL)
    if subtitle:
        text(s, subtitle, .48, sep + .14, 9.0, .7, 13, GRAY)


def card(s, x, y, w, h, tag, title, body):
    rect(s, x, y, w, h, WHITE, BORDER, True)
    text(s, tag, x + .16, y + .14, w - .32, .25, 11, TEAL, True)
    text(s, title, x + .16, y + .49, w - .32, .56, 17, NAVY, True)
    text(s, body, x + .16, y + 1.11, w - .32, h - 1.2, 12.5, GRAY)


def evidence(title, image, bullets):
    s = slide()
    head(s, "SECClaw live product evidence", title, title_size=27)
    rect(s, .48, 1.35, 4.32, 3.96, WHITE, BORDER, True)
    s.shapes.add_picture(str(ASSETS / image), Inches(.58), Inches(1.45), width=Inches(4.12), height=Inches(3.76))
    for i, bullet in enumerate(bullets):
        y = 1.62 + i * 1.04
        text(s, "●", 5.18, y, .27, .32, 15, TEAL, True)
        text(s, bullet, 5.47, y, 4.0, .8, 15, NAVY)
    rect(s, 5.26, 4.96, 4.14, .39, RGBColor(226, 246, 247), RGBColor(226, 246, 247), True)
    text(s, "LOCAL FIRST  •  POLICY GOVERNED", 5.52, 5.04, 3.8, .22, 10.5, TEAL, True)
    return s


s = slide()
text(s, "SECURITY, WITH A SAFETY CATCH", .5, .44, 7, .32, 12, TEAL, True)
text(s, "SecClaw", .5, .86, 6.2, .72, 40, NAVY, True)
text(s, "Local security operations that can explain, ask, act, and prove.", .5, 1.78, 6.1, 1.45, 25, NAVY, True)
text(s, "Dell × NVIDIA AI Hackathon  •  San Francisco  •  2026", .5, 3.55, 6.4, .4, 13, GRAY)
rect(s, .5, 4.37, 6.0, .8, WHITE, BORDER, True)
text(s, "Built for enterprise IT and SecOps teams\nAll inference runs locally on the GB10", .72, 4.51, 5.6, .55, 16, NAVY, True)
rect(s, 7.05, .75, 2.45, 4.4, RGBColor(226, 245, 240), RGBColor(180, 222, 213), True)
text(s, "GB10", 7.37, 1.41, 1.9, .73, 38, TEAL, True)
text(s, "Local AI.\nHuman authority.\nVerifiable results.", 7.37, 2.42, 1.85, 1.8, 20, NAVY, True)

s = slide()
head(s, "THE PROBLEM", "Security teams are stuck between speed and control.",
     "Sensitive infrastructure data should stay on site. Remediation needs explicit authorization.", 27)
card(s, .48, 2.39, 2.89, 2.72, "01", "Blind spots", "Small IT teams cannot continuously inspect every internal device and exposed service.")
card(s, 3.56, 2.39, 2.89, 2.72, "02", "Cloud exposure", "Topology, versions, findings, and conversations are too sensitive for remote LLM workflows.")
card(s, 6.64, 2.39, 2.89, 2.72, "03", "Unsafe automation", "A plausible suggestion is not permission to change production state.")

evidence("Unified Security Dashboard", "dashboard-screenshot.jpg", [
    "19 unique devices consolidated in PostgreSQL",
    "Live exposure and scan coverage charts",
    "One inventory shared across every workflow",
])

s = slide()
head(s, "THE ASK", "Let SecClaw watch the risk while humans keep the authority.", title_size=27)
text(s, "Built today on Dell Pro Max with GB10", .5, 2.17, 8.9, .45, 21, NAVY, True)
text(s, "gpt-oss-120b  •  vLLM  •  OpenShell  •  Next.js  •  PostgreSQL", .5, 2.81, 8.9, .43, 17, GRAY)
rect(s, .5, 3.82, 8.9, .94, WHITE, BORDER, True)
text(s, "github.com/cupertinose/SecClaw", .79, 4.12, 8.3, .42, 20, TEAL, True)

s = slide()
head(s, "THE SOLUTION", "Discover. Explain. Approve. Act. Verify.",
     "SecClaw turns observations into controlled action, with local reasoning on GB10.", 28)
steps = [
    ("1", "DISCOVER", "Fixed, scoped scans and mDNS inventory"),
    ("2", "EXPLAIN", "Evidence grounded local risk analysis"),
    ("3", "APPROVE", "Human decision bound to a plan hash"),
    ("4", "ACT", "OpenShell sandbox with explicit policy"),
    ("5", "VERIFY", "Rescan results and preserve an audit trail"),
]
for i, (num, label, body) in enumerate(steps):
    x = .48 + i * 1.91
    rect(s, x, 2.4, 1.76, 2.65, WHITE, BORDER, True)
    text(s, num, x + .13, 2.59, 1.5, .4, 20, TEAL, True)
    text(s, label, x + .13, 3.12, 1.52, .36, 13.4, NAVY, True)
    text(s, body, x + .13, 3.63, 1.49, 1.14, 11.4, GRAY)

s = slide()
head(s, "ARCHITECTURE", "Everything important stays on the Dell Pro Max with GB10",
     "128 GB unified memory supports the local model and full application stack.", 25)
architecture = [
    ("SECCLAW DASHBOARD", "Next.js  •  PostgreSQL  •  structured audit timeline"),
    ("AGENT + POLICY", "Scoped tools  •  approval gate  •  plan hashing  •  rollback"),
    ("LOCAL INFERENCE", "gpt-oss-120b via vLLM  •  no cloud LLM calls"),
    ("GUARDED EXECUTION", "OpenShell sandbox  •  read-only inventory  •  network policy"),
]
for i, (label, detail) in enumerate(architecture):
    y = 2.27 + i * .78
    rect(s, .55, y, 8.9, .61, RGBColor(220, 244, 237) if i == 2 else WHITE, BORDER, True)
    text(s, label, .75, y + .15, 2.3, .31, 13.5, NAVY, True)
    text(s, detail, 3.22, y + .15, 5.94, .33, 12.8, GRAY)

s = slide()
head(s, "LIVE DEMO", "From exposed service to verified remediation",
     "The walkthrough shows the working system, including the refusal path.", 27)
demo = [
    ("Scope", "Admin authorizes a host, ports, time window, and policy."),
    ("Detect", "SecClaw finds an exposed service and preserves evidence."),
    ("Reason", "Local gpt-oss-120b proposes safe validation."),
    ("Approve", "A human reviews the change summary and plan hash."),
    ("Act + verify", "OpenShell acts under policy; a rescan confirms the outcome."),
]
for i, (label, detail) in enumerate(demo):
    y = 2.05 + i * .65
    rect(s, .53, y, .51, .48, RGBColor(220, 244, 237), TEAL, True)
    text(s, str(i + 1), .71, y + .065, .25, .3, 15, TEAL, True)
    text(s, label, 1.25, y + .08, 1.5, .35, 15, NAVY, True)
    text(s, detail, 2.78, y + .08, 6.55, .45, 13.3, GRAY)

evidence("Auditable Scan History", "scans-screenshot.jpg", [
    "Manual, imported, and scheduled scans",
    "Responding hosts and service evidence",
    "Every completed result persists to PostgreSQL",
])

s = slide()
head(s, "BUSINESS VALUE", "A security analyst in a box, with humans in charge.",
     "Reduce time to understand and remediate while preserving accountability.", 27)
card(s, .48, 2.39, 2.89, 2.72, "LOCAL", "Private by design", "Infrastructure evidence and inference remain on premises.")
card(s, 3.56, 2.39, 2.89, 2.72, "SAFE", "Approval gated", "Every state change is scoped, time bound, and auditable.")
card(s, 6.64, 2.39, 2.89, 2.72, "FAST", "One workflow", "Discovery, explanation, verification, and rollback live in one place.")

evidence("Device Discovery & Actions", "devices-screenshot.jpg", [
    "mDNS hostnames and advertised services",
    "19 device unified inventory",
    "One click remediation and SSH reporting",
])

evidence("OpenShell Governed SSH", "openshell-screenshot.jpg", [
    "Gateway connected with mTLS transport",
    "Sandbox ready under accepted policy v2",
    "Configured SSH target: 192.168.43.161",
])

P.save(OUTPUT)
print(OUTPUT)
